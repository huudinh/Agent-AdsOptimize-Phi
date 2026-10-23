# -*- coding: utf-8 -*-
"""
Sinh file Excel "Bộ từ khoá Google Ads" cho 1 BỘ VÙNG x DỊCH VỤ.

Generator BRAND-NEUTRAL: tên thương hiệu đọc từ khoá zone.thuong_hieu trong file JSON,
không hardcode. Dùng chung được cho mọi gói anh em trong Project/Agent.

Chia theo VÙNG (bán kính đi lại) x DỊCH VỤ, không chia theo ZONE dịch vụ.
    Tầng thị trường -> Chân dung KH -> Hành trình S1..S7 -> Cụm truy vấn ưu tiên
    -> Landing (L1..L6) -> Từ khoá (có market + geo_mode) -> Phủ định -> Ngưỡng

Dùng:
    python tools/build_keyword_workbook.py markets/rhinoplasty-tier-c-au.json
    python tools/build_keyword_workbook.py "markets/*.json" -o out/

Yêu cầu: pip install openpyxl

Năm sheet xuất ra:
    1. Tổng quan          - thị trường, chân dung KH, phân bổ ngân sách, phễu, nguyên tắc
    2. Bộ từ khoá         - 1 dòng = 1 từ khoá x kiểu khớp, CÓ market + geo_mode + WIN /4
    3. Từ khoá phủ định   - danh sách chung + phủ định chéo điều hướng truy vấn
    4. Hành trình KH      - S1..S7, landing page, ngân hàng Hook x Format
    5. Giá thầu & Tối ưu  - chấm điểm ưu tiên, ngưỡng, lộ trình 3 pha, A/B test, rủi ro

BA CỔNG CHẶN BUILD (xem validate()):
    - cổng CLAIM   : từ cấm tiếng Anh + claim đã gỡ (lifetime warranty, board certified...)
    - cổng GEO-TARGET     : từ khoá có tên thành phố VN mà thiếu market / geo_mode
    - cổng REGULATED: market rủi ro cao (AU/UK/EU/GCC/US) mà ghi_chu chưa khai cách
                      xử lý creative rủi ro (before-after / testimonial kết quả)
"""
import argparse
import glob
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Console Windows mặc định là cp1252 -> thông báo lỗi tiếng Việt bị hỏng chữ.
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

# ---------------------------------------------------------------- bảng màu
NAVY = "FF005F88"   # --teal của Design System International v1.0
BAND = "FFDBE5F2"
ZEBRA = "FFF4F7FC"
WHITE = "FFFFFFFF"
TOTAL = "FFF2F2F2"
INPUT_BG = "FFFFF2CC"
INPUT_FG = "FF0000FF"
FORMULA_FG = "FF007F00"

FONT = "Arial"
# Định dạng số tiền: KHÔNG kèm ký hiệu tiền tệ, nên dùng được cho VND, PHP,
# USD… Đơn vị khai ở zone.don_vi_tien và in ra đầu sheet 1.
VND = "#,##0"
PCT = "0%"
PCT1 = "0.0%"
NUM0 = "#,##0"
NUM2 = "0.00"
THIN = Side(style="thin", color="FFD9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _title(ws, row, text, size=14):
    c = ws.cell(row, 1, text)
    c.font = Font(name=FONT, size=size, bold=True, color=NAVY)
    return row + 1


def _note(ws, row, text):
    c = ws.cell(row, 1, text)
    c.font = Font(name=FONT, size=10, italic=True, color="FF667085")
    return row + 1


def _header(ws, row, labels):
    for i, lb in enumerate(labels, start=1):
        c = ws.cell(row, i, lb)
        c.font = Font(name=FONT, size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    ws.row_dimensions[row].height = 30
    return row + 1


def _band(ws, row, text, span=6):
    c = ws.cell(row, 1, text)
    c.font = Font(name=FONT, size=10, bold=True, color=NAVY)
    for i in range(1, span + 1):
        ws.cell(row, i).fill = PatternFill("solid", fgColor=BAND)
    return row + 1


def _cell(ws, row, col, value, fmt=None, kind=None, bold=False,
          fill=None, align=None, wrap=True):
    """kind: None | 'input' (nền vàng, chữ xanh) | 'formula' (chữ xanh lá)"""
    c = ws.cell(row, col, value)
    color = None
    if kind == "input":
        c.fill = PatternFill("solid", fgColor=INPUT_BG)
        color = INPUT_FG
    elif kind == "formula":
        color = FORMULA_FG
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.font = Font(name=FONT, size=10, bold=bold, color=color)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    c.border = BORDER
    return c


def _widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ============================================================ 1. Tổng quan
def sheet_tong_quan(wb, z):
    ws = wb.create_sheet("1. Tổng quan")
    _widths(ws, {"A": 34, "B": 20, "C": 46, "D": 34, "E": 22, "F": 22, "G": 12})
    zone = z["zone"]
    r = 1

    r = _title(ws, r, "BỘ TỪ KHOÁ GOOGLE ADS: " + zone["ten"].upper()
                  + " | " + str(zone.get("thuong_hieu", "[CHUA KHAI thuong_hieu]")).upper())
    r = _note(ws, r, "Mục tiêu: " + zone["muc_tieu"])
    r = _note(ws, r, "Đơn vị tiền tệ: " + str(zone.get("don_vi_tien", "[CHUA KHAI don_vi_tien]")))
    r = _note(ws, r, "Thị trường: " + str(zone.get("markets", "[CHUA KHAI]"))
                     + "  |  Tầng: " + str(zone.get("tang", "[CHUA KHAI]"))
                     + "  |  Ngôn ngữ: " + str(zone.get("ngon_ngu", "[CHUA KHAI]"))
                     + "  |  Rủi ro quảng cáo: " + str(zone.get("rui_ro_qc", "[CHUA KHAI]")))
    r = _note(ws, r, "Phương pháp: VÙNG x DỊCH VỤ -> chân dung KH -> hành trình theo chặng "
                     "-> cụm truy vấn ưu tiên -> landing -> từ khoá (có market + geo_mode).")
    r = _note(ws, r, "Cổng GEO-TARGET: mọi dòng từ khoá phải khai market (người gõ ĐANG Ở ĐÂU, không phải "
                     "thành phố trong từ khoá) và geo_mode. Phòng khám phục vụ khách ĐẾN TẬN NƠI "
                     "thì MỌI dòng bắt buộc geo_mode = Presence.")
    r = _note(ws, r, "Ô nền vàng chữ xanh là giả định đầu vào, thay bằng số thực của phòng khám. "
                     "Ô chữ xanh lá là công thức, không sửa tay.")
    r += 1

    # ---- A. chân dung khách hàng
    r = _title(ws, r, "A. CHÂN DUNG KHÁCH HÀNG - GỐC CỦA MỌI QUYẾT ĐỊNH PHÍA SAU", 12)
    r = _note(ws, r, "Chọn VÙNG trước, chọn NGƯỜI, rồi mới chọn từ khoá. Tách khách lần đầu và khách đi sửa lại.")
    r = _header(ws, r, ["Mã", "Chân dung KH", "Insight / nỗi đau thật", "Rào cản chốt chính",
                        "Vào phễu ở chặng", "Chiến dịch phục vụ", "Giá trị ca"])
    for p in z["chan_dung"]:
        _cell(ws, r, 1, p["ma"], bold=True, align="center")
        _cell(ws, r, 2, p["ten"], bold=True)
        _cell(ws, r, 3, p["insight"])
        _cell(ws, r, 4, p["rao_can"])
        _cell(ws, r, 5, p["chang_vao"], align="center")
        _cell(ws, r, 6, p["chien_dich"])
        _cell(ws, r, 7, p["gia_tri_ca"], align="center")
        r += 1
    r += 1

    # ---- B. phân bổ ngân sách
    r = _title(ws, r, "B. PHÂN BỔ NGÂN SÁCH THEO CHIẾN DỊCH", 12)
    _cell(ws, r, 1, "Ngân sách Search/tháng", bold=True)
    budget_ref = "$B$" + str(r)
    _cell(ws, r, 2, zone["ngan_sach_thang"], fmt=VND, kind="input", align="center")
    _cell(ws, r, 3, "Giả định. Thay bằng ngân sách được duyệt.")
    r += 2

    r = _header(ws, r, ["Chiến dịch", "Ưu tiên", "Vai trò trong hành trình KH", "Tỷ trọng ngân sách",
                        "Ngân sách/tháng", "Ngân sách/ngày", "Số từ khoá"])
    cd_first = r
    cd_rows = []
    for cd in z["chien_dich"]:
        _cell(ws, r, 1, cd["ten"], bold=True)
        _cell(ws, r, 2, cd["uu_tien"], align="center")
        _cell(ws, r, 3, cd["vai_tro"])
        _cell(ws, r, 4, cd["ty_trong"], fmt=PCT, kind="input", align="center")
        _cell(ws, r, 5, "=" + budget_ref + "*D" + str(r), fmt=VND, kind="formula")
        _cell(ws, r, 6, "=E" + str(r) + "/30", fmt=VND, kind="formula")
        _cell(ws, r, 7, "=COUNTIF('2. Bộ từ khoá'!$B:$B,A" + str(r) + ")",
              kind="formula", align="center")
        cd_rows.append(r)
        r += 1
    cd_last = r - 1

    _cell(ws, r, 1, "Tổng", bold=True, fill=TOTAL)
    _cell(ws, r, 2, None, fill=TOTAL)
    _cell(ws, r, 3, None, fill=TOTAL)
    _cell(ws, r, 4, "=SUM(D%d:D%d)" % (cd_first, cd_last), fmt=PCT, bold=True, fill=TOTAL, align="center")
    _cell(ws, r, 5, "=SUM(E%d:E%d)" % (cd_first, cd_last), fmt=VND, bold=True, fill=TOTAL)
    _cell(ws, r, 6, "=SUM(F%d:F%d)" % (cd_first, cd_last), fmt=VND, bold=True, fill=TOTAL)
    _cell(ws, r, 7, "=SUM(G%d:G%d)" % (cd_first, cd_last), bold=True, fill=TOTAL, align="center")
    r += 2

    r = _header(ws, r, ["Cấp ưu tiên", "Tỷ trọng ngân sách", "Ngân sách/tháng", "Quy tắc đấu thầu"])
    for pr in z["cap_uu_tien"]:
        _cell(ws, r, 1, pr["cap"], bold=True, align="center")
        _cell(ws, r, 2, "=SUMIF($B$%d:$B$%d,A%d,$D$%d:$D$%d)" % (cd_first, cd_last, r, cd_first, cd_last),
              fmt=PCT, kind="formula", align="center")
        _cell(ws, r, 3, "=SUMIF($B$%d:$B$%d,A%d,$E$%d:$E$%d)" % (cd_first, cd_last, r, cd_first, cd_last),
              fmt=VND, kind="formula")
        _cell(ws, r, 4, pr["quy_tac"])
        r += 1
    r += 1

    # ---- C. mô hình phễu
    r = _title(ws, r, "C. MÔ HÌNH PHỄU: TỪ NGÂN SÁCH ĐẾN DOANH THU", 12)
    r = _header(ws, r, ["Chỉ số", "Giá trị", "Ghi chú / nguồn số"])
    p = z["phieu"]
    # Phễu quốc tế 8 bước, gom về 5 mốc đo được:
    #   click -> lead -> ĐẶT CỌC -> ĐẾN KHÁM -> ca chốt
    # Đặt cọc là điểm chuyển trạng thái thật của du lịch y tế: khách đã đặt cọc
    # thì gần như luôn đến. Xem kn-chan-doan-chi-so-ph.md §1.
    inputs = [("cpc", "CPC trung bình", VND),
              ("click_lead", "Tỷ lệ click -> lead", PCT1),
              ("lead_lich", "Tỷ lệ lead -> ĐẶT CỌC (sau khi phản hồi được)", PCT1),
              ("show_up", "Tỷ lệ đặt cọc -> ĐẾN KHÁM (arrival)", PCT1),
              ("chot", "Tỷ lệ đến khám -> ca chốt", PCT1),
              ("gia_tri_ca", "Giá trị ca trung bình", VND)]
    ref = {}
    for key, label, fmt in inputs:
        _cell(ws, r, 1, label, bold=True)
        _cell(ws, r, 2, p[key], fmt=fmt, kind="input", align="center")
        _cell(ws, r, 3, p.get("ghi_chu", {}).get(key, ""))
        ref[key] = "B" + str(r)
        r += 1
    r += 1

    def outrow(label, formula, fmt, note):
        nonlocal r
        _cell(ws, r, 1, label, bold=True)
        _cell(ws, r, 2, formula, fmt=fmt, kind="formula", align="center")
        _cell(ws, r, 3, note)
        cur = "B" + str(r)
        r += 1
        return cur

    click = outrow("Click", "=IF(%s=0,0,%s/%s)" % (ref["cpc"], budget_ref, ref["cpc"]), NUM0, "Công thức")
    lead = outrow("Lead", "=%s*%s" % (click, ref["click_lead"]), NUM0,
                  "Lead = form + chat WhatsApp/Zalo có số + cuộc gọi >= 60 giây")
    lich = outrow("Đặt cọc", "=%s*%s" % (lead, ref["lead_lich"]), NUM0,
                  "Điểm chuyển trạng thái thật của du lịch y tế")
    kham = outrow("Đến khám (arrival)", "=%s*%s" % (lich, ref["show_up"]), NUM0, "Công thức")
    chot = outrow("Ca chốt", "=%s*%s" % (kham, ref["chot"]), NUM0, "Công thức")
    doanh_thu = outrow("Doanh thu", "=%s*%s" % (chot, ref["gia_tri_ca"]), VND, "Công thức")
    cpl = outrow("CPL", "=IF(%s=0,0,%s/%s)" % (lead, budget_ref, lead), VND,
                 "CẨN TRỌNG: so CPL giữa hai market là so hai thứ khác nhau")
    cp_arrival = outrow("Chi phí / khách ĐẾN KHÁM", "=IF(%s=0,0,%s/%s)" % (kham, budget_ref, kham),
                        VND, "CHỈ SỐ BẮC ĐẨU của gói quốc tế")
    cp_case = outrow("Chi phí / ca chốt", "=IF(%s=0,0,%s/%s)" % (chot, budget_ref, chot), VND,
                     "Chỉ số quyết định ngân sách, không phải CPL")
    outrow("ROAS", "=IF(%s=0,0,%s/%s)" % (budget_ref, doanh_thu, budget_ref), NUM2, "Công thức")
    r += 1

    # ---- D. nguyên tắc vận hành
    r = _title(ws, r, "D. NGUYÊN TẮC VẬN HÀNH", 12)
    for i, rule in enumerate(z["nguyen_tac"], start=1):
        _cell(ws, r, 1, "%d. %s" % (i, rule), wrap=False)
        r += 1

    return {"cd_rows": cd_rows, "cd_first": cd_first, "cd_last": cd_last,
            "cpl": cpl, "cp_case": cp_case, "cp_arrival": cp_arrival,
            "gia_tri_ca": ref["gia_tri_ca"],
            "show_up": ref["show_up"], "chot": ref["chot"], "lead_lich": ref["lead_lich"]}


# ========================================================== 2. Bộ từ khoá
KW_COLS = [
    ("stt", "STT", 6),
    ("chien_dich", "Chiến dịch", 34),
    ("nhom_qc", "Nhóm quảng cáo", 24),
    ("tu_khoa", "Từ khoá", 40),
    ("khop", "Kiểu khớp", 11),
    ("market", "Market (người gõ ở đâu)", 13),
    ("geo_mode", "Geo mode", 17),
    ("uu_tien", "Ưu tiên", 9),
    ("chang", "Chặng (S1-S7)", 13),
    ("chan_dung", "Chân dung KH", 12),
    ("rao_can", "Rào cản chốt", 13),
    ("canh_tranh", "Cạnh tranh dự kiến", 13),
    ("win", "Cổng WIN (/4)", 10),
    ("lp", "Landing page", 12),
    ("thong_diep", "Thông điệp / CTA chính", 40),
    ("volume", "Lượng tìm kiếm/tháng (Keyword Planner - theo market)", 16),
    ("cpc", "CPC đề xuất (Keyword Planner - theo market)", 16),
    ("ghi_chu", "Ghi chú vận hành", 46),
]
CENTER_KEYS = {"stt", "khop", "market", "geo_mode", "uu_tien", "chang", "chan_dung",
               "canh_tranh", "win", "lp", "volume", "cpc"}


def sheet_tu_khoa(wb, z):
    ws = wb.create_sheet("2. Bộ từ khoá")
    _widths(ws, dict((get_column_letter(i), w) for i, (_, _, w) in enumerate(KW_COLS, start=1)))
    _header(ws, 1, [lb for _, lb, _ in KW_COLS])

    r = 2
    for i, kw in enumerate(z["tu_khoa"], start=1):
        bg = WHITE if i % 2 else ZEBRA
        row = dict(kw)
        row["stt"] = i
        for ci, (key, _, _) in enumerate(KW_COLS, start=1):
            v = row.get(key, "")
            _cell(ws, r, ci, v if v != "" else None, fill=bg,
                  align="center" if key in CENTER_KEYS else None,
                  fmt=None,
                  bold=(key == "tu_khoa"))
        r += 1

    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(KW_COLS)), r - 1)
    ws.freeze_panes = "E2"
    return r - 2


# ==================================================== 3. Từ khoá phủ định
def sheet_phu_dinh(wb, z):
    ws = wb.create_sheet("3. Từ khoá phủ định")
    _widths(ws, {"A": 34, "B": 30, "C": 12, "D": 46, "E": 50})
    r = _header(ws, 1, ["Nhóm", "Từ khoá phủ định", "Kiểu khớp", "Áp dụng ở", "Lý do"])

    r = _band(ws, r, "A. DANH SÁCH CHUNG - cấp tài khoản: " + z["zone"]["ds_phu_dinh_chung"], 5)
    for n in z["phu_dinh_chung"]:
        _cell(ws, r, 1, n["nhom"])
        _cell(ws, r, 2, n["tu"], bold=True)
        _cell(ws, r, 3, n["khop"], align="center")
        _cell(ws, r, 4, n["ap_dung"])
        _cell(ws, r, 5, n["ly_do"])
        r += 1
    r += 1

    r = _band(ws, r, "B. PHỦ ĐỊNH CHÉO - điều hướng mỗi truy vấn về đúng 1 chiến dịch", 5)
    for n in z["phu_dinh_cheo"]:
        _cell(ws, r, 1, n["nhom"])
        _cell(ws, r, 2, n["tu"], bold=True)
        _cell(ws, r, 3, n["khop"], align="center")
        _cell(ws, r, 4, n["ap_dung"])
        _cell(ws, r, 5, n["ly_do"])
        r += 1
    r += 1

    r = _band(ws, r, "C. QUY TRÌNH CẬP NHẬT", 5)
    for line in z["quy_trinh_phu_dinh"]:
        _cell(ws, r, 1, line, wrap=False)
        r += 1

    ws.freeze_panes = "A2"


# ====================================================== 4. Hành trình KH
JOURNEY_COLS = [
    ("chang", "Chặng (S1-S7)", 26),
    ("chan_dung", "Chân dung KH chính", 20),
    ("tam_ly", "Tâm lý & insight", 38),
    ("rao_can", "Rào cản chốt", 18),
    ("chien_dich", "Chiến dịch", 30),
    ("tu_khoa", "Cụm truy vấn đại diện", 30),
    ("thong_diep", "Thông điệp quảng cáo", 38),
    ("bang_chung", "Bằng chứng bắt buộc", 34),
    ("lp", "Landing page", 14),
    ("chuyen_doi", "Chuyển đổi đo lường", 26),
    ("sau_lead", "Hành động sau lead (CRM)", 34),
    ("kpi", "KPI chặng", 22),
    ("tiep_theo", "Bước tiếp theo", 26),
]


def sheet_hanh_trinh(wb, z):
    ws = wb.create_sheet("4. Hành trình KH")
    _widths(ws, dict((get_column_letter(i), w) for i, (_, _, w) in enumerate(JOURNEY_COLS, start=1)))
    r = _header(ws, 1, [lb for _, lb, _ in JOURNEY_COLS])

    for i, s in enumerate(z["hanh_trinh"]):
        bg = WHITE if i % 2 == 0 else ZEBRA
        for ci, (key, _, _) in enumerate(JOURNEY_COLS, start=1):
            _cell(ws, r, ci, s.get(key) or None, fill=bg,
                  bold=(key == "chang"),
                  align="center" if key in ("lp", "chan_dung") else None)
        r += 1
    r += 1

    r = _band(ws, r, "B. LANDING PAGE - 1 nhóm = 1 dịch vụ x 1 chặng x 1 intent = 1 trang", 8)
    r = _header(ws, r, ["Landing page", "Loại khung", "Chặng", "Chân dung KH",
                        "Rào cản gỡ chính", "Nội dung bắt buộc", "CTA chính",
                        "Micro-conversion trước form"])
    for lp in z["landing"]:
        _cell(ws, r, 1, lp["ma_ten"], bold=True)
        _cell(ws, r, 2, lp.get("loai", ""), align="center")
        _cell(ws, r, 3, lp.get("chang", ""), align="center")
        _cell(ws, r, 4, lp.get("chan_dung", ""), align="center")
        _cell(ws, r, 5, lp.get("rao_can", ""))
        _cell(ws, r, 6, lp["noi_dung"])
        _cell(ws, r, 7, lp["cta"])
        _cell(ws, r, 8, lp.get("micro", ""))
        r += 1
    r += 1

    if z.get("hook_format"):
        r = _band(ws, r, "C. NGÂN HÀNG HOOK x FORMAT (FB / TikTok) - cùng chân dung, cùng chặng với Search", 5)
        r = _header(ws, r, ["Chặng", "Chân dung KH", "Hook ưu tiên", "Format khả thi - hiệu quả",
                            "Mục tiêu thống nhất với Search"])
        for h in z["hook_format"]:
            _cell(ws, r, 1, h["chang"], bold=True, align="center")
            _cell(ws, r, 2, h["chan_dung"], align="center")
            _cell(ws, r, 3, h["hook"])
            _cell(ws, r, 4, h["format"])
            _cell(ws, r, 5, h["muc_tieu"])
            r += 1

    ws.freeze_panes = "B2"


# ================================================= 5. Giá thầu & Tối ưu
def sheet_gia_thau(wb, z, a):
    ws = wb.create_sheet("5. Giá thầu & Tối ưu")
    _widths(ws, {"A": 36, "B": 17, "C": 18, "D": 18, "E": 14, "F": 11, "G": 36, "H": 36})
    r = 1
    r = _title(ws, r, "A. CHẤM ĐIỂM ƯU TIÊN CHIẾN DỊCH", 12)
    r = _note(ws, r, "Điểm = Ý định mua x Giá trị ca x Khả năng chốt (thang 1-5). "
                     "Điểm ban đầu là đánh giá chuyên môn; cập nhật hằng tháng bằng tỷ lệ đến khám "
                     "thực tế từ CRM. Lưu ý độ trễ 4-12 tuần: không kết luận tầng B/C "
                     "trong cửa sổ dưới 8 tuần.")
    r = _header(ws, r, ["Chiến dịch", "Ý định mua (1-5)", "Giá trị ca (1-5)", "Khả năng chốt (1-5)",
                        "Điểm ưu tiên", "Xếp hạng", "Chiến lược giá thầu khởi đầu", "Chỉ tiêu cạnh tranh"])
    first = r
    last = first + len(z["chien_dich"]) - 1
    for cd, src in zip(z["chien_dich"], a["cd_rows"]):
        d = cd["diem"]
        _cell(ws, r, 1, "='1. Tổng quan'!A" + str(src), kind="formula", bold=True)
        _cell(ws, r, 2, d["y_dinh"], kind="input", align="center")
        _cell(ws, r, 3, d["gia_tri"], kind="input", align="center")
        _cell(ws, r, 4, d["kha_nang_chot"], kind="input", align="center")
        _cell(ws, r, 5, "=B%d*C%d*D%d" % (r, r, r), kind="formula", align="center")
        _cell(ws, r, 6, "=RANK(E%d,$E$%d:$E$%d)" % (r, first, last), kind="formula", align="center")
        _cell(ws, r, 7, cd["bid"])
        _cell(ws, r, 8, cd["chi_tieu"])
        r += 1
    r += 1

    r = _title(ws, r, "B. NGƯỠNG QUẢN TRỊ GIÁ THẦU", 12)
    r = _header(ws, r, ["Chỉ số", "Giá trị", "Ghi chú"])
    _cell(ws, r, 1, "CPL mục tiêu", bold=True)
    _cell(ws, r, 2, "='1. Tổng quan'!" + a["cpl"], fmt=VND, kind="formula", align="center")
    _cell(ws, r, 3, "Lấy từ mô hình phễu tab 1")
    cpl_row = r
    r += 1
    _cell(ws, r, 1, "Chi phí / khách ĐẾN KHÁM mục tiêu", bold=True)
    _cell(ws, r, 2, "='1. Tổng quan'!" + a["cp_arrival"], fmt=VND, kind="formula", align="center")
    _cell(ws, r, 3, "CHỈ SỐ BẮC ĐẨU - đánh giá chiến dịch bằng cột này, không bằng CPL")
    r += 1
    _cell(ws, r, 1, "Chi phí / ca chốt mục tiêu", bold=True)
    _cell(ws, r, 2, "='1. Tổng quan'!" + a["cp_case"], fmt=VND, kind="formula", align="center")
    _cell(ws, r, 3, "Lấy từ mô hình phễu tab 1")
    case_row = r
    r += 1
    _cell(ws, r, 1, "Giá trị gán cho 1 lượt đặt cọc", bold=True)
    _cell(ws, r, 2, "='1. Tổng quan'!%s*'1. Tổng quan'!%s*'1. Tổng quan'!%s"
          % (a["gia_tri_ca"], a["show_up"], a["chot"]), fmt=VND, kind="formula", align="center")
    _cell(ws, r, 3, "Dùng khi import chuyển đổi offline (Pha 3)")
    lich_row = r
    r += 1
    _cell(ws, r, 1, "Giá trị gán cho 1 lead", bold=True)
    _cell(ws, r, 2, "=B%d*'1. Tổng quan'!%s" % (lich_row, a["lead_lich"]),
          fmt=VND, kind="formula", align="center")
    _cell(ws, r, 3, "Dùng khi import chuyển đổi offline (Pha 3)")
    r += 2

    r = _header(ws, r, ["Quy tắc", "Hệ số", "Ngưỡng", "Hành động"])
    for rule in z["nguong"]:
        _cell(ws, r, 1, rule["quy_tac"], bold=True)
        _cell(ws, r, 2, rule["he_so"], kind="input", align="center")
        base = "$B$" + str(case_row if rule["moc"] == "ca_chot" else cpl_row)
        _cell(ws, r, 3, "=B%d*%s" % (r, base), fmt=VND, kind="formula", align="center")
        _cell(ws, r, 4, rule["hanh_dong"])
        r += 1
    r += 1

    r = _title(ws, r, "C. LỘ TRÌNH CHIẾN LƯỢC GIÁ THẦU 3 PHA", 12)
    r = _header(ws, r, ["Pha", "Thời gian", "Điều kiện chuyển pha", "Chiến lược giá thầu",
                        "Chuyển đổi tối ưu", "Việc bắt buộc"])
    for p in z["lo_trinh"]:
        _cell(ws, r, 1, p["pha"], bold=True)
        _cell(ws, r, 2, p["thoi_gian"], align="center")
        _cell(ws, r, 3, p["dieu_kien"])
        _cell(ws, r, 4, p["chien_luoc"])
        _cell(ws, r, 5, p["chuyen_doi"])
        _cell(ws, r, 6, p["bat_buoc"])
        r += 1
    r += 1

    r = _title(ws, r, "D. QUY TẮC CẠNH TRANH & VẬN HÀNH", 12)
    r = _header(ws, r, ["Hạng mục", "Quy tắc"])
    for q in z["quy_tac_canh_tranh"]:
        _cell(ws, r, 1, q["hang_muc"], bold=True)
        _cell(ws, r, 2, q["quy_tac"])
        r += 1
    r += 1

    r = _title(ws, r, "E. KẾ HOẠCH A/B TEST", 12)
    r = _header(ws, r, ["Test", "Biến thể A", "Biến thể B", "Chỉ số quyết định", "Thời gian tối thiểu"])
    for t in z["ab_test"]:
        _cell(ws, r, 1, t["test"], bold=True)
        _cell(ws, r, 2, t["a"])
        _cell(ws, r, 3, t["b"])
        _cell(ws, r, 4, t["chi_so"])
        _cell(ws, r, 5, t["thoi_gian"], align="center")
        r += 1
    r += 1

    r = _title(ws, r, "F. RỦI RO VÀ CÁCH KIỂM SOÁT", 12)
    r = _header(ws, r, ["Rủi ro", "Cách kiểm soát"])
    for k in z["rui_ro"]:
        _cell(ws, r, 1, k["rui_ro"], bold=True)
        _cell(ws, r, 2, k["kiem_soat"])
        r += 1


# ================================================================ kiểm tra
REQUIRED = ["zone", "chan_dung", "chien_dich", "cap_uu_tien", "phieu", "nguyen_tac",
            "tu_khoa", "phu_dinh_chung", "phu_dinh_cheo", "quy_trinh_phu_dinh",
            "hanh_trinh", "landing", "nguong", "lo_trinh", "quy_tac_canh_tranh",
            "ab_test", "rui_ro"]

# Từ cấm — bản tiếng Anh (tầng A/C/D) VÀ tiếng Việt (tầng B kiều bào).
# Bảng đầy đủ + từ thay thế: knowledge/kn-rao-phap-ly-ph.md §1
CAM = [
    # English
    "best ", " best", "the best", "no.1", "no. 1", "#1", "top-rated", "top rated",
    "leading ", "finest", "premier ", "guaranteed", "guarantee result",
    "permanent", "forever", "painless", "pain-free", "pain free",
    "no scar", "scar-free", "scar free", "risk-free", "risk free",
    "completely safe", "absolutely safe", "no complication", "100% success",
    "cheapest", "unbeatable", "lowest price", "miracle", "flawless",
    # Vietnamese (tầng B)
    "tốt nhất", "số 1", "số một", "không đau 100", "cam kết thành công",
    "khỏi 100", "đẹp tuyệt đối", "vĩnh viễn", "duy nhất",
    "an toàn tuyệt đối", "không biến chứng", "không để lại sẹo",
]


# ------------------------------------------------------- cổng CLAIM ĐÃ GỠ
# BRAND-NEUTRAL: danh sách dưới đây chỉ giữ những claim rủi ro ở MỌI thị trường.
# Claim rủi ro riêng theo thương hiệu / theo nước khai ở `zone.claim_cam` trong file
# JSON, dạng [{"tu": "...", "ly_do": "..."}]. Ví dụ: gói Kangnam Quốc tế chặn
# "board certified" vì nhầm với tư cách hội; gói Philippines KHÔNG chặn vì bác sĩ
# có chứng nhận thật (Diplomate PDS) — đó là tri thức của module thương hiệu,
# không phải của generator.
CLAIM_GO_MAC_DINH = [
    (u"lifetime warranty", u"cam kết trọn đời không có T&C công bố -> không dùng. "
                           u"Viết 'follow-up and warranty policy - see terms'"),
    (u"lifetime guarantee", u"như trên"),
    (u"bảo hành vĩnh viễn", u"như trên"),
    (u"guaranteed result", u"cam kết kết quả điều trị là vùng cấm của quảng cáo y tế"),
    (u"guaranteed results", u"như trên"),
    (u"100% safe", u"cam kết tuyệt đối - không chứng minh được"),
    (u"no side effects", u"cam kết tuyệt đối - không chứng minh được"),
    (u"painless", u"cam kết tuyệt đối - không chứng minh được"),
]
# Câu tự nhắc đã đúng luật thì không báo lỗi (vd ghi_chu nói rõ "không dùng lifetime").
MIEN_TRU = u"không dùng lifetime"
CLAIM_GO = list(CLAIM_GO_MAC_DINH)


def nap_claim_cam(zone):
    """Gộp claim cấm riêng của thương hiệu (zone.claim_cam) vào danh sách mặc định."""
    global CLAIM_GO
    CLAIM_GO = list(CLAIM_GO_MAC_DINH)
    for it in (zone.get("claim_cam") or []):
        tu = (it.get("tu") or "").strip().lower()
        if tu:
            CLAIM_GO.append((tu, it.get("ly_do", u"claim cấm theo module thương hiệu")))


def soat_claim(chuoi, nhan, errs):
    """Thấy claim đã gỡ trong một ô nội dung thì ghi lỗi."""
    low = (chuoi or "").lower()
    if MIEN_TRU in low:
        return
    for w, vi_sao in CLAIM_GO:
        if w in low:
            errs.append(u"%s: còn claim đã gỡ '%s' - %s" % (nhan, w, vi_sao))

# ------------------------------------------------------------------ cổng GEO-TARGET
# Cùng một từ khoá là HAI người khác nhau:
#   "plastic surgery hanoi" gõ từ Sydney  = khách tầng C sắp bay sang
#   "plastic surgery hanoi" gõ từ Hà Nội  = expat tầng D đang sống ở đó
# Một landing, một mẫu QC -> mất cả hai. Xem kn-vung-philippines.md §2.
# Để Google tự chọn kiểu nhắm cho chiến dịch expat -> tràn traffic quốc tế,
# CPL trông rẻ mà không ai đến khám, và lỗi này IM LẶNG.
TP_VN = ["hanoi", "ha noi", "hà nội", "ho chi minh", "hochiminh", "hcmc", "saigon",
         "sai gon", "sài gòn", "da nang", "đà nẵng", "vietnam", "viet nam", "việt nam"]
GEO_MODE_OK = ["presence", "presence+interest", "presence + interest",
               "presence or interest"]

# ------------------------------------------------------------ cổng REGULATED
# Market rủi ro cao: creative dùng before-after / testimonial kết quả lâm sàng
# phải khai cách xử lý trong ghi_chu. AHPRA (Úc) cấm thẳng testimonial kết quả.
MARKET_RUI_RO_CAO = ["AU", "UK", "GB", "NZ", "US", "CA", "IE",
                     "DE", "FR", "NL", "SE", "CZ", "EU", "EEA",
                     "AE", "SA", "KW", "QA", "BH", "OM", "GCC"]
# Ghi chú được coi là đã xử lý khi nêu một trong các cụm sau.
REGULATED_OK = ["no before-after", "khong before-after", "không before-after",
                "no testimonial", "không testimonial", "khong testimonial",
                "non-clinical", "phi lâm sàng", "phi lam sang",
                "ahpra", "asa", "cap code", "ftc", "gdpr",
                "pháp chế", "phap che", "legal review", "chờ duyệt", "cho duyet"]


def validate(z, path):
    nap_claim_cam(z.get("zone", {}))
    errs = []
    for k in REQUIRED:
        if k not in z:
            errs.append("thiếu khoá bắt buộc: " + k)
    if errs:
        raise SystemExit("[LỖI] " + path + "\n  - " + "\n  - ".join(errs))

    tong = round(sum(c["ty_trong"] for c in z["chien_dich"]), 6)
    if abs(tong - 1.0) > 1e-6:
        errs.append("tổng tỷ trọng ngân sách = %.4f, phải bằng 1.0" % tong)

    ten_cd = set(c["ten"] for c in z["chien_dich"])
    ma_lp = set(lp["ma_ten"].split(":")[0].strip() for lp in z["landing"])
    ma_cd = set(p["ma"] for p in z["chan_dung"])
    cap = set(p["cap"] for p in z["cap_uu_tien"])

    for c in z["chien_dich"]:
        if c["uu_tien"] not in cap:
            errs.append("chiến dịch '%s': cấp ưu tiên '%s' chưa khai báo ở cap_uu_tien"
                        % (c["ten"], c["uu_tien"]))

    seen = set()
    for i, kw in enumerate(z["tu_khoa"], start=1):
        tag = "từ khoá #%d '%s'" % (i, kw.get("tu_khoa", "?"))
        if kw.get("chien_dich") not in ten_cd:
            errs.append("%s: chiến dịch '%s' không có trong danh sách chiến dịch"
                        % (tag, kw.get("chien_dich")))
        if kw.get("lp") and kw["lp"] not in ma_lp:
            errs.append("%s: landing '%s' chưa khai báo ở mục landing" % (tag, kw["lp"]))
        if kw.get("chan_dung") and kw["chan_dung"] not in ma_cd:
            errs.append("%s: chân dung '%s' chưa khai báo ở mục chan_dung" % (tag, kw["chan_dung"]))
        if kw.get("khop") == "Broad":
            errs.append("%s: Broad match bị cấm ở Pha 1 - xem nguyên tắc vận hành" % tag)
        low = (kw.get("thong_diep", "") or "").lower()
        for w in CAM:
            if w in low:
                errs.append("%s: thông điệp chứa từ cấm quảng cáo y tế '%s'" % (tag, w))
        # Cổng CLAIM chỉ quét những gì CHÚNG TA viết: thông điệp + tên nhóm QC.
        # KHÔNG quét `tu_khoa`: từ khoá là chữ KHÁCH gõ, không phải claim của mình.
        # "cheaper than korea" là truy vấn thật và được phép nhắm; điều bị cấm là
        # viết lại nó thành claim trong thông điệp. Xem kn-rao-phap-ly-ph.md §4 luật ③.
        soat_claim(u"%s %s" % (kw.get("thong_diep", ""), kw.get("nhom_qc", "")), tag, errs)
        # Cùng một chuỗi từ khoá ở HAI market là HAI người khác nhau, không phải
        # trùng lặp: "how long does it last" gõ từ V1 và từ V2 cần bid và landing
        # khác nhau. Vì vậy khoá trùng lặp phải gồm cả market.
        key = ((kw.get("tu_khoa") or "").strip().lower(), kw.get("khop"),
               (kw.get("market") or "").strip().upper())
        if key in seen:
            errs.append("trùng từ khoá x kiểu khớp x market: %s (%s, %s)" % key)
        seen.add(key)

        # ---- cổng GEO-TARGET: mọi dòng phải khai market + geo_mode ----
        mk = (kw.get("market") or "").strip()
        gm = (kw.get("geo_mode") or "").strip()
        if not mk:
            errs.append("%s: thiếu 'market' - không biết người gõ ĐANG Ở ĐÂU thì không "
                        "biết luật quảng cáo nào áp, viết ngôn ngữ nào, gắn kênh nào "
                        "(kn-vung-philippines.md §2)" % tag)
        if not gm:
            errs.append("%s: thiếu 'geo_mode' - phải khai Presence hoặc Presence+Interest. "
                        "Để Google tự chọn cho chiến dịch expat là lỗi IM LẶNG: tràn "
                        "traffic quốc tế, CPL rẻ giả, không ai đến khám" % tag)
        elif gm.lower() not in GEO_MODE_OK:
            errs.append("%s: geo_mode '%s' không hợp lệ - dùng 'Presence' hoặc "
                        "'Presence+Interest'" % (tag, gm))
        # Tầng D (expat tại VN) bắt buộc Presence, nếu không sẽ hút traffic du lịch y tế
        if mk.upper() == "VN-EXPAT" and gm.lower() not in ("presence",):
            errs.append("%s: market VN-EXPAT (expat đang sống tại VN) bắt buộc "
                        "geo_mode='Presence'. Dùng Presence+Interest ở đây sẽ hút toàn bộ "
                        "traffic du lịch y tế vào landing sai" % tag)
        # Từ khoá có tên thành phố/quốc gia VN là bẫy hai-người-một-từ-khoá
        kwl = (kw.get("tu_khoa") or "").lower()
        if any(c in kwl for c in TP_VN) and (not mk or not gm):
            errs.append("%s: từ khoá có địa danh Việt Nam nên là BẪY GEO (khách sắp bay "
                        "vs expat đang ở đó) - bắt buộc khai cả market và geo_mode" % tag)

        # ---- cổng REGULATED: market rủi ro cao phải khai cách xử lý creative ----
        if mk and mk.split("-")[0].upper() in MARKET_RUI_RO_CAO:
            ghi = (kw.get("ghi_chu", "") or "").lower()
            if not any(w in ghi for w in REGULATED_OK):
                errs.append(
                    "%s: market '%s' thuộc nhóm RỦI RO CAO (AHPRA/ASA/FTC/GDPR/GCC) mà "
                    "ghi_chu chưa khai cách xử lý creative. Nêu rõ một trong: landing "
                    "KHÔNG dùng before-after / KHÔNG dùng testimonial kết quả lâm sàng / "
                    "chỉ dùng testimonial phi lâm sàng / đang chờ Pháp chế thị trường duyệt. "
                    "Xem kn-rao-phap-ly-ph.md §2." % (tag, mk))

    for i, ht in enumerate(z["hanh_trinh"], start=1):
        for k in ("thong_diep", "bang_chung", "tiep_theo", "chien_dich"):
            soat_claim(ht.get(k, ""), u"hành trình #%d · %s" % (i, k), errs)
    for i, lp in enumerate(z["landing"], start=1):
        for k in ("ma_ten", "noi_dung", "cta"):
            soat_claim(lp.get(k, ""), u"landing #%d · %s" % (i, k), errs)
    for c in z["chien_dich"]:
        soat_claim(c.get("ten", ""), u"chiến dịch '%s'" % c.get("ten", ""), errs)

    if errs:
        raise SystemExit("[LỖI] " + path + "\n  - " + "\n  - ".join(errs))


def build(path, outdir):
    with open(path, encoding="utf-8") as f:
        z = json.load(f)
    validate(z, path)

    wb = Workbook()
    wb.remove(wb.active)
    anchors = sheet_tong_quan(wb, z)
    n_kw = sheet_tu_khoa(wb, z)
    sheet_phu_dinh(wb, z)
    sheet_hanh_trinh(wb, z)
    sheet_gia_thau(wb, z, anchors)
    wb.active = 0

    name = z["zone"].get("ten_file") or ("KN _ Google Ads _ Bộ từ khoá " + z["zone"]["ten"])
    os.makedirs(outdir, exist_ok=True)
    dest = os.path.join(outdir, name + ".xlsx")
    wb.save(dest)
    return dest, n_kw, len(z["chien_dich"]), len(z["chan_dung"])


def main():
    ap = argparse.ArgumentParser(
        description="Sinh bộ từ khoá Google Ads theo VÙNG x DỊCH VỤ (brand-neutral)")
    ap.add_argument("configs", nargs="+", help="file JSON cấu hình ZONE (hỗ trợ wildcard)")
    ap.add_argument("-o", "--outdir", default="out", help="thư mục xuất (mặc định: out/)")
    args = ap.parse_args()

    paths = []
    for c in args.configs:
        paths.extend(sorted(glob.glob(c)) or [c])

    for p in paths:
        dest, n_kw, n_cd, n_persona = build(p, args.outdir)
        print("[OK] %s\n     -> %s\n     %d chân dung KH | %d chiến dịch | %d từ khoá"
              % (p, dest, n_persona, n_cd, n_kw))


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:
        pass
    main()
